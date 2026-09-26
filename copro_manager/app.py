from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date, timedelta
import io
import os
import json
from sqlalchemy.orm import joinedload
from database import COPROPRIETES_DATA, db, Copropriete, Civilite # <-- Ajoute COPROPRIETES_DATA
from honoraires_data import HONORAIRE_DATA
from visites_data import ANNEES as VISITES_ANNEES, LIGNES as VISITES_LIGNES
from travaux_data import LIGNES as TRAVAUX_LIGNES
from assurance_data import ASSURANCE_COLONNES, ASSURANCE_LIGNES
from coproprietaires_data import COPROPRIETAIRES_EXCEL_DATA
import json as _json_module

# ========== CONFIGURATION ==========
app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY') or os.urandom(24)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///copro_manager.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = 'static/uploads'
db = SQLAlchemy(app)

# ===== CONTEXT PROCESSOR (pour rendre Civilite accessible dans les templates) =====
@app.context_processor
def inject_models():
    return dict(Civilite=Civilite)

# Ensure upload folder exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# ========== MODELS ==========
class Copropriete(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    numero = db.Column(db.Integer, unique=True, nullable=False)
    nom = db.Column(db.String(200), nullable=False)
    date_mise_copropriete = db.Column(db.Date)
    programme_neolia = db.Column(db.String(100))
    adresse = db.Column(db.String(300))
    ville = db.Column(db.String(100))
    immatriculation = db.Column(db.String(100))
    nombre_logements = db.Column(db.Integer)
    exercice_comptable = db.Column(db.String(20))
    gestionnaire = db.Column(db.String(50))
    comptable = db.Column(db.String(100))
    responsable_secteur = db.Column(db.String(200))
    commercial = db.Column(db.String(200))
    notaire = db.Column(db.String(200))
    est_active = db.Column(db.Boolean, default=True)

    # Relationships
    fiche_immeuble = db.relationship('FicheImmeuble', backref='copropriete', uselist=False)
    contrats = db.relationship('Contrat', backref='copropriete', lazy=True)
    coproprietaires = db.relationship('Coproprietaire', backref='copropriete', lazy=True)
    assemblees_generales = db.relationship('AssembleeGenerale', backref='copropriete', lazy=True)
    resolutions_futures = db.relationship('ResolutionFuture', backref='copropriete', lazy=True)
    demandes = db.relationship('Demande', backref='copropriete', lazy=True, cascade="all, delete-orphan")

class FicheImmeuble(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    copropriete_id = db.Column(db.Integer, db.ForeignKey('copropriete.id'), unique=True, nullable=False)
    date_construction = db.Column(db.Date)
    date_arrete_compte = db.Column(db.Date)
    reference_neolia = db.Column(db.String(100))
    nom_commercial_vente_hlm = db.Column(db.String(200))
    nom_responsable_secteur = db.Column(db.String(200))
    nom_notaire = db.Column(db.String(200))
    designation_immeuble = db.Column(db.String(300))
    iban = db.Column(db.String(100))
    ics = db.Column(db.String(50))
    lien_dossier = db.Column(db.String(500))
    lien_budget = db.Column(db.String(500))

class Contrat(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    copropriete_id = db.Column(db.Integer, db.ForeignKey('copropriete.id'), nullable=False)
    type_contrat = db.Column(db.String(100))
    nature = db.Column(db.String(200))
    fournisseur = db.Column(db.String(200))
    date_debut = db.Column(db.Date)
    date_fin = db.Column(db.Date)
    montant_annuel = db.Column(db.Float)
    prestations = db.relationship('Prestation', backref='contrat', lazy=True, cascade="all, delete-orphan")

class Prestation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    contrat_id = db.Column(db.Integer, db.ForeignKey('contrat.id'), nullable=False)
    libelle = db.Column(db.String(200))
    frequence = db.Column(db.String(50))
    prix_unitaire = db.Column(db.Float)
    quantite = db.Column(db.Integer)
    total = db.Column(db.Float)

class Coproprietaire(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    civilite_id = db.Column(db.Integer, db.ForeignKey('civilites.id'), nullable=True)
    civilite = db.relationship('Civilite', backref='coproprietaires', foreign_keys=[civilite_id])
    copropriete_id = db.Column(db.Integer, db.ForeignKey('copropriete.id'), nullable=False)
    date_acquisition = db.Column(db.Date)
    nom = db.Column(db.String(100))
    prenom = db.Column(db.String(100))
    numero_lot = db.Column(db.String(50))
    nature_lot = db.Column(db.String(100))
    email = db.Column(db.String(200))
    telephone = db.Column(db.String(50))
    est_residence_principale = db.Column(db.Boolean, default=False)
    est_loue = db.Column(db.Boolean, default=False)
    est_conseil_syndical = db.Column(db.Boolean, default=False)
    date_envoi_mail_accueil = db.Column(db.Date)
    lien_espace_client = db.Column(db.String(500))
    # Adresse postale du copropriétaire (pour correspondre au fichier Excel
    # « Copropriétaires.xlsx » et à l'onglet Copropriétaires).
    adresse = db.Column(db.String(300))
    code_postal = db.Column(db.String(20))
    ville = db.Column(db.String(100))

    locataire_nom = db.Column(db.String(100))
    locataire_civilite_id = db.Column(db.Integer, db.ForeignKey('civilites.id'), nullable=True)
    locataire_prenom = db.Column(db.String(100))
    locataire_email = db.Column(db.String(200))
    locataire_telephone = db.Column(db.String(50))

    lots = db.relationship('LotCoproprietaire', backref='coproprietaire', lazy=True, cascade="all, delete-orphan")
    emails = db.relationship('EmailCoproprietaire', backref='coproprietaire', lazy=True, cascade="all, delete-orphan")
    telephones = db.relationship('TelephoneCoproprietaire', backref='coproprietaire', lazy=True, cascade="all, delete-orphan")

    @property
    def nom_affiche(self):
        parts = [p for p in (self.nom, self.prenom) if p]
        return ' '.join(parts) or f'Copropriétaire #{self.id}'

class Civilite(db.Model):
    __tablename__ = 'civilites'
    id = db.Column(db.Integer, primary_key=True)
    libelle = db.Column(db.String(20), unique=True, nullable=False)

class LotCoproprietaire(db.Model):
    __tablename__ = 'lots_coproprietaires'
    id = db.Column(db.Integer, primary_key=True)
    numero = db.Column(db.String(50), nullable=False)
    nature = db.Column(db.String(100))
    coproprietaire_id = db.Column(db.Integer, db.ForeignKey('coproprietaire.id'), nullable=False)

class EmailCoproprietaire(db.Model):
    __tablename__ = 'emails_coproprietaires'
    id = db.Column(db.Integer, primary_key=True)
    value = db.Column(db.String(200), nullable=False)
    principal = db.Column(db.Boolean, default=False)
    coproprietaire_id = db.Column(db.Integer, db.ForeignKey('coproprietaire.id'), nullable=False)

class TelephoneCoproprietaire(db.Model):
    __tablename__ = 'telephones_coproprietaires'
    id = db.Column(db.Integer, primary_key=True)
    value = db.Column(db.String(50), nullable=False)
    principal = db.Column(db.Boolean, default=False)
    coproprietaire_id = db.Column(db.Integer, db.ForeignKey('coproprietaire.id'), nullable=False)

class AssembleeGenerale(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    copropriete_id = db.Column(db.Integer, db.ForeignKey('copropriete.id'), nullable=False)
    date = db.Column(db.Date)
    horaire_debut = db.Column(db.String(20))
    lieu = db.Column(db.String(200))
    type_ag = db.Column(db.String(100))  # Nouveau champ: Type d'AG

    points_a_retenir = db.relationship('PointARetenir', backref='assemblee_generale', lazy=True, cascade="all, delete-orphan")
    budgets_travaux = db.relationship('BudgetTravaux', backref='assemblee_generale', lazy=True, cascade="all, delete-orphan")
class PointARetenir(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ag_id = db.Column(db.Integer, db.ForeignKey('assemblee_generale.id'), nullable=False)
    description = db.Column(db.String(500))

class BudgetTravaux(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ag_id = db.Column(db.Integer, db.ForeignKey('assemblee_generale.id'), nullable=False)
    nom_travaux = db.Column(db.String(200))
    entreprise_retenue = db.Column(db.String(200))
    montant_travaux = db.Column(db.Float)
    montant_honoraires_syndic = db.Column(db.Float)
    total_budget_vote = db.Column(db.Float)
    utilise_fond_travaux = db.Column(db.Boolean, default=False)
    montant_fond_travaux = db.Column(db.Float)
    nombre_appels = db.Column(db.Integer)
    cle_repartition = db.Column(db.String(100))
    appels_fonds = db.relationship('AppelFonds', backref='budget_travaux', lazy=True, cascade="all, delete-orphan")

class AppelFonds(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    budget_travaux_id = db.Column(db.Integer, db.ForeignKey('budget_travaux.id'), nullable=False)
    date_appel = db.Column(db.Date)
    montant_exige_pourcentage = db.Column(db.Float)

class ResolutionFuture(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    copropriete_id = db.Column(db.Integer, db.ForeignKey('copropriete.id'), nullable=False)
    titre = db.Column(db.String(200))
    projet = db.Column(db.String(500))

class Demande(db.Model):
    __tablename__ = 'demandes'
    id = db.Column(db.Integer, primary_key=True)
    copropriete_id = db.Column(db.Integer, db.ForeignKey('copropriete.id'), nullable=False)
    date = db.Column(db.Date)
    demandeur = db.Column(db.String(200))
    type_action = db.Column(db.String(50))
    details = db.Column(db.Text)
    echeance = db.Column(db.String(200))
    suivi = db.Column(db.String(20), default='a_faire')  # a_faire | en_cours | fait

# ========== MODÈLES POUR LE TABLEAU ASSURANCE (dynamique) ==========
# Approche entité-attribut-valeur pour permettre l'ajout/suppression
# libre de lignes (copropriétés) et de colonnes (années/indicateurs).

class AssuranceColonne(db.Model):
    __tablename__ = 'assurance_colonnes'
    id = db.Column(db.Integer, primary_key=True)
    groupe = db.Column(db.String(50), nullable=False, default='Tarif')
    en_tete = db.Column(db.String(200), nullable=False)
    ordre = db.Column(db.Integer, nullable=False, default=0)
    # formule: type de calcul automatique, vide si colonne saisie à la main.
    #   prix_m2        : valeur = cotisation_annee / superficie
    #   evolution_n   : valeur = (cotisation_N - cotisation_N-1) / cotisation_N-1
    #   evolution_base: valeur = (cotisation_N - cotisation_base) / cotisation_base
    formule = db.Column(db.String(50), nullable=True)
    # parametres formule (JSON) : ex {"annee":12,"ref":11} pour l'année/colonne de référence
    formule_params = db.Column(db.Text, nullable=True)
    cellules = db.relationship('AssuranceCellule', backref='colonne', lazy=True, cascade="all, delete-orphan")

class AssuranceLigne(db.Model):
    __tablename__ = 'assurance_lignes'
    id = db.Column(db.Integer, primary_key=True)
    ordre = db.Column(db.Integer, nullable=False, default=0)
    est_total = db.Column(db.Boolean, default=False)
    cellules = db.relationship('AssuranceCellule', backref='ligne', lazy=True, cascade="all, delete-orphan")

class AssuranceCellule(db.Model):
    __tablename__ = 'assurance_cellules'
    id = db.Column(db.Integer, primary_key=True)
    ligne_id = db.Column(db.Integer, db.ForeignKey('assurance_lignes.id'), nullable=False)
    colonne_id = db.Column(db.Integer, db.ForeignKey('assurance_colonnes.id'), nullable=False)
    valeur = db.Column(db.Text)

# ========== MODÈLE POUR LE TABLEAU HONORAIRES ("Honoraire agate") ==========
# Une ligne = une copropriété. Les colonnes "€ TTC 1er et 2e Trimestre",
# "€ TTC 3e et 4e Trimestre", "Total € TTC" et "Total € HT" sont calculées
# automatiquement (recalculées à chaque sauvegarde), sauf pour les lignes
# dont le type de facturation est "Spécifique" (montants figés).

TYPE_FACTURATION = ['Classique', 'Forfait', 'ASL', 'Spécifique']

class HonoraireLigne(db.Model):
    __tablename__ = 'honoraire_lignes'
    id = db.Column(db.Integer, primary_key=True)
    numero = db.Column(db.Integer)  # N° de copro
    ordre = db.Column(db.Integer, nullable=False, default=0)
    mise_copro = db.Column(db.Date)          # Mise en copro
    fin_contrat = db.Column(db.Date)         # Fin de contrat
    type_facturation = db.Column(db.String(50))  # Type de facturation
    logements = db.Column(db.Integer)         # Logements
    tarif_s1 = db.Column(db.Float)            # € TTC / LP 1er semestre
    tarif_s2 = db.Column(db.Float)            # € TTC / LP 2e semestre
    t12 = db.Column(db.Float)                 # € TTC 1er et 2e Trimestre
    t34 = db.Column(db.Float)                 # € TTC 3e et 4e Trimestre
    total_ttc = db.Column(db.Float)           # Total € TTC
    total_ht = db.Column(db.Float)            # Total € HT
    est_total = db.Column(db.Boolean, default=False)

    @property
    def est_specifique(self):
        return (self.type_facturation or '').strip().lower() == 'spécifique'

# ========== MODÈLES POUR LE TABLEAU VISITES D'IMMEUBLE (dynamique) ==========
# Approche identique au tableau Assurance : colonnes = années de visite,
# lignes = copropriétés, cellules = date de visite (ou texte libre).
# L'ajout/suppression libre de colonnes (années) et de lignes (copros) est
# possible, et chaque cellule est éditable depuis le tableau ou la fiche copro.

class VisiteColonne(db.Model):
    __tablename__ = 'visite_colonnes'
    id = db.Column(db.Integer, primary_key=True)
    annee = db.Column(db.String(10), nullable=False)  # en-tête: année de visite
    ordre = db.Column(db.Integer, nullable=False, default=0)
    cellules = db.relationship('VisiteCellule', backref='colonne', lazy=True, cascade="all, delete-orphan")

class VisiteLigne(db.Model):
    __tablename__ = 'visite_lignes'
    id = db.Column(db.Integer, primary_key=True)
    numero = db.Column(db.Integer)  # N° de copro
    mise_copro = db.Column(db.String(20))  # date de mise en copro (JJ/MM/AAAA)
    ordre = db.Column(db.Integer, nullable=False, default=0)
    cellules = db.relationship('VisiteCellule', backref='ligne', lazy=True, cascade="all, delete-orphan")

class VisiteCellule(db.Model):
    __tablename__ = 'visite_cellules'
    id = db.Column(db.Integer, primary_key=True)
    ligne_id = db.Column(db.Integer, db.ForeignKey('visite_lignes.id'), nullable=False)
    colonne_id = db.Column(db.Integer, db.ForeignKey('visite_colonnes.id'), nullable=False)
    valeur = db.Column(db.String(100))  # date de visite (JJ/MM/AAAA) ou texte libre


# ========== MODÈLES POUR LE TABLEAU TRAVAUX (dynamique) ==========

STATUTS_TRAVAUX = ['en_cours', 'pret_ag', 'termine']

class TravauxLigne(db.Model):
    """Une ligne du tableau « ADF Travaux » : un travaux pour une copropriété."""
    __tablename__ = 'travaux_lignes'
    id = db.Column(db.Integer, primary_key=True)
    numero = db.Column(db.Integer)             # N° de copro
    fin_exercice = db.Column(db.String(20))     # « 30 juin » ou « 31 décembre »
    date_ag = db.Column(db.Date)                # tri : la plus récente en premier
    type = db.Column(db.String(200))            # description du travaux
    montant = db.Column(db.Float)               # €
    honoraire = db.Column(db.Float)             # €
    honoraire_facture = db.Column(db.Boolean, default=False)
    nb_appels = db.Column(db.Integer)          # nb d'appels de fonds
    tethrawin = db.Column(db.String(200))       # texte libre
    devis_valide = db.Column(db.String(100))    # date ou texte libre
    os = db.Column(db.String(200))              # texte libre
    facture = db.Column(db.String(200))        # texte libre
    annee_cloture = db.Column(db.String(20))    # année (ex: 2027) ou texte
    statut = db.Column(db.String(20), default='en_cours')  # en_cours/pret_ag/termine
    fait = db.Column(db.Boolean, default=False)  # suivi comptable : appel de fonds fait
    appels = db.relationship('TravauxAppel', backref='ligne', lazy=True,
                             cascade="all, delete-orphan",
                             order_by='TravauxAppel.ordre')

class TravauxAppel(db.Model):
    """Une date d'appel de fonds d'une ligne de travaux."""
    __tablename__ = 'travaux_appels'
    id = db.Column(db.Integer, primary_key=True)
    ligne_id = db.Column(db.Integer, db.ForeignKey('travaux_lignes.id'), nullable=False)
    ordre = db.Column(db.Integer, nullable=False, default=0)
    date_appel = db.Column(db.String(20))       # JJ/MM/AAAA ou texte libre
    fait = db.Column(db.Boolean, default=False)  # suivi comptable : appel de fonds fait

# ========== UTILITY FUNCTIONS ==========
def parse_date(date_str):
    """Parse date in DD/MM/YYYY or YYYY-MM-DD format"""
    if not date_str:
        return None
    try:
        if '/' in date_str:
            return datetime.strptime(date_str, '%d/%m/%Y').date()
        return datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None

# ========== STATISTICS FUNCTION ==========
def get_statistiques():
    from sqlalchemy import func
    total_copros = Copropriete.query.count()
    total_active = Copropriete.query.filter_by(est_active=True).count()
    total_inactive = total_copros - total_active
    total_logements = db.session.query(func.coalesce(func.sum(Copropriete.nombre_logements), 0)).scalar() or 0

    stats_gestionnaires = db.session.query(
        Copropriete.gestionnaire,
        func.count(Copropriete.id).label('count'),
        func.sum(Copropriete.nombre_logements).label('logements')
    ).filter(
        Copropriete.gestionnaire.isnot(None),
        Copropriete.est_active == True
    ).group_by(Copropriete.gestionnaire).all()

    gestionnaires = {}
    for g in stats_gestionnaires:
        gestionnaires[g.gestionnaire] = {
            'count': g.count,
            'logements': g.logements,
            'percentage': round((g.count / total_active) * 100, 1) if total_active > 0 else 0
        }

    stats_periodes = db.session.query(
        Copropriete.exercice_comptable,
        func.count(Copropriete.id).label('count')
    ).filter(
        Copropriete.exercice_comptable.isnot(None),
        Copropriete.est_active == True
    ).group_by(Copropriete.exercice_comptable).all()

    periodes = {}
    for p in stats_periodes:
        periodes[p.exercice_comptable] = {
            'count': p.count,
            'percentage': round((p.count / total_active) * 100, 1) if total_active > 0 else 0
        }

    return {
        'total_logements': total_logements,
        'gestionnaires': gestionnaires,
        'periodes': periodes,
        'total_active': total_active,
        'total_inactive': total_inactive
    }

# ========== ROUTES ==========
@app.route('/contrats')
def contrats():
    types_contrats = [
        "Assurance",
        "Nettoyage",
        "Espaces verts",
        "Sécurité incendie",
        "Sous-compteurs d'eau",
        "VMC",
        "Ascenseurs",
        "Électricité",
        "Eau",
        "Chaufferie",
        "Gaz",
        "Portes automatiques"
    ]
    
    descriptions = {
        "Assurance": "Contrats d'assurance pour les copropriétés",
        "Nettoyage": "Contrats de nettoyage des parties communes",
        "Espaces verts": "Contrats d'entretien des espaces verts",
        "Sécurité incendie": "Contrats de maintenance des systèmes de sécurité incendie",
        "Sous-compteurs d'eau": "Contrats de gestion des sous-compteurs d'eau",
        "VMC": "Contrats de maintenance des ventilations mécaniques contrôlées",
        "Ascenseurs": "Contrats de maintenance des ascenseurs",
        "Électricité": "Contrats d'approvisionnement et maintenance électrique",
        "Eau": "Contrats d'approvisionnement en eau",
        "Chaufferie": "Contrats de maintenance des systèmes de chauffage",
        "Gaz": "Contrats d'approvisionnement en gaz",
        "Portes automatiques": "Contrats de maintenance des portes automatiques"
    }
    
    types_contrats.sort()
    return render_template('contrats.html', types_contrats=types_contrats, descriptions=descriptions)

@app.route('/contrats/<type_contrat>')
def type_contrat_page(type_contrat):
    # Page dédiée pour le suivi des contrats d'assurance
    if type_contrat == 'assurance':
        return redirect(url_for('assurance_page'))
    # Page générique : à terme, un template par type pourra être ajouté
    flash("Cette page n'est pas encore disponible pour ce type de contrat.", 'info')
    return redirect(url_for('contrats'))


def _importer_assurance():
    """Importe le tableau Assurance dans la base (une seule fois, si les
    tables sont vides). Les données viennent de l'ancien fichier
    « Assurance.xlsx », intégrées au logiciel : plus besoin du fichier."""
    if AssuranceColonne.query.count() > 0 or AssuranceLigne.query.count() > 0:
        return
    colonnes = []
    for c in ASSURANCE_COLONNES:
        col = AssuranceColonne(
            groupe=c['groupe'],
            en_tete=c['en_tete'],
            ordre=ASSURANCE_COLONNES.index(c),
            formule=c['formule'],
            formule_params=_json_module.dumps(c['formule_params']) if c['formule_params'] else None,
        )
        db.session.add(col)
        colonnes.append(col)
    db.session.flush()
    ordre_ligne = 0
    for l in ASSURANCE_LIGNES:
        ligne = AssuranceLigne(ordre=ordre_ligne, est_total=l['est_total'])
        db.session.add(ligne)
        db.session.flush()
        ordre_ligne += 1
        for idx, valeur in l['valeurs'].items():
            if int(idx) < len(colonnes):
                db.session.add(AssuranceCellule(
                    ligne_id=ligne.id, colonne_id=colonnes[int(idx)].id,
                    valeur=valeur))
    db.session.commit()


def _assurance_context():
    """Construit le contexte de la page Assurance depuis la base."""
    colonnes = AssuranceColonne.query.order_by(AssuranceColonne.ordre).all()
    lignes = AssuranceLigne.query.order_by(AssuranceLigne.ordre).all()

    # group_headers: nom du groupe pour chaque colonne (pour le colspan)
    group_headers = [c.groupe for c in colonnes]
    sub_headers = [c.en_tete for c in colonnes]

    # Construction d'une matrice de cellules: rows[ligne][colonne]
    rows = []
    for ligne in lignes:
        cellules = {c.colonne_id: c.valeur for c in ligne.cellules}
        row = {
            'id': ligne.id,
            'est_total': ligne.est_total,
            'valeurs': [cellules.get(col.id, '') for col in colonnes],
        }
        rows.append(row)

    total_row = next((r for r in rows if r['est_total']), None)
    data_rows = [r for r in rows if not r['est_total']]

    # Map id->en-tête pour l'édition de cellule
    colonnes_map = {c.id: c.en_tete for c in colonnes}

    # Données des formules pour le recalcul JS temps réel
    import json as _json
    colonnes_formules = {}
    for c in colonnes:
        if c.formule:
            colonnes_formules[c.id] = {
                'formule': c.formule,
                'params': _json.loads(c.formule_params) if c.formule_params else {},
            }

    # Index des colonnes par id pour le JS
    colonnes_ordre = {c.id: i for i, c in enumerate(colonnes)}

    return {
        'group_headers': group_headers,
        'sub_headers': sub_headers,
        'rows': data_rows,
        'total_row': total_row,
        'colonnes': colonnes,
        'colonnes_map': colonnes_map,
        'colonnes_formules_json': _json.dumps(colonnes_formules),
        'colonnes_ordre_json': _json.dumps(colonnes_ordre),
    }


@app.route('/contrats/assurance', endpoint='assurance_page')
def assurance_page():
    """Affiche le tableau de suivi tarifaire des contrats d'assurance,
    chargé depuis la base de données (modèle dynamique)."""
    _importer_assurance()
    return render_template('assurance.html', **_assurance_context())


@app.route('/contrats/assurance/export', endpoint='assurance_export')
def assurance_export():
    """Exporte le tableau d'assurance dans un fichier Excel (.xlsx).

    Reconstruit le tableau tel qu'affiché : une ligne d'en-têtes de groupe
    (Copropriété / Contrat / Tarif) fusionnée, une ligne d'en-têtes de
    colonnes, puis une ligne par copropriété (et la ligne TOTAL en dernier).
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    colonnes = AssuranceColonne.query.order_by(AssuranceColonne.ordre).all()
    lignes = AssuranceLigne.query.order_by(AssuranceLigne.ordre).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assurance"

    # Styles
    gras_blanc = Font(bold=True, color="FFFFFF")
    fond_groupe = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    fond_sous = PatternFill(start_color="495057", end_color="495057", fill_type="solid")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    nb_colonnes = len(colonnes)
    if nb_colonnes == 0:
        flash("Aucune donnée à exporter.", "error")
        return redirect(url_for('assurance_page'))

    # Ligne 1 : en-têtes de groupe (fusion des colonnes d'un même groupe)
    idx = 0
    while idx < nb_colonnes:
        groupe = colonnes[idx].groupe
        debut = idx
        while idx < nb_colonnes and colonnes[idx].groupe == groupe:
            idx += 1
        ws.cell(row=1, column=debut + 1, value=groupe)
        if idx - debut > 1:
            ws.merge_cells(start_row=1, start_column=debut + 1,
                           end_row=1, end_column=idx)
        cell = ws.cell(row=1, column=debut + 1)
        cell.font = gras_blanc
        cell.fill = fond_groupe
        cell.alignment = centre
        cell.border = bordure

    # Ligne 2 : en-têtes de colonnes
    for c, col in enumerate(colonnes):
        cell = ws.cell(row=2, column=c + 1, value=col.en_tete)
        cell.font = gras_blanc
        cell.fill = fond_sous
        cell.alignment = centre
        cell.border = bordure

    # Lignes de données
    ligne_excel = 3
    for ligne in lignes:
        cellules = {cl.colonne_id: cl.valeur for cl in ligne.cellules}
        for c, col in enumerate(colonnes):
            valeur = cellules.get(col.id, '')
            if valeur == '' or valeur is None:
                continue
            # Tente de convertir en nombre quand c'est possible
            try:
                nombre = float(str(valeur).replace(' ', '').replace(',', '.'))
                if nombre == int(nombre):
                    nombre = int(nombre)
            except (ValueError, TypeError):
                nombre = str(valeur)
            cell = ws.cell(row=ligne_excel, column=c + 1, value=nombre)
            cell.border = bordure
        # Ligne TOTAL en gras
        if ligne.est_total:
            for c in range(nb_colonnes):
                cell = ws.cell(row=ligne_excel, column=c + 1)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        ligne_excel += 1

    # Largeurs de colonnes automatiques (approximation sur l'en-tête)
    for c, col in enumerate(colonnes):
        longueur = max(len(col.en_tete.split('\n')[0]) if col.en_tete else 0, 10)
        ws.column_dimensions[openpyxl.utils.get_column_letter(c + 1)].width = longueur + 2
    # Première colonne un peu plus large (n° de copro)
    if colonnes:
        ws.column_dimensions['A'].width = 14
    ws.freeze_panes = "B3"

    # Nom du fichier : assurance_AAAA-MM-JJ.xlsx
    nom_fichier = f"assurance_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    return send_file(
        tampon,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _recalculer_ligne(ligne):
    """Recalcule toutes les cellules formules d'une ligne d'assurance.
    À appeler après modification d'une cotisation ou de la superficie."""
    import json as _json
    colonnes = AssuranceColonne.query.order_by(AssuranceColonne.ordre).all()
    # Les paramètres de formule stockent des INDEX (0-N), pas des IDs.
    # On construit la correspondance index -> id de colonne.
    idx_vers_id = {i: c.id for i, c in enumerate(colonnes)}
    cellules = {c.colonne_id: c for c in ligne.cellules}

    def valeur_num(idx):
        col_id = idx_vers_id.get(idx)
        if col_id is None:
            return None
        cell = cellules.get(col_id)
        if not cell or not cell.valeur:
            return None
        try:
            v = float(str(cell.valeur).replace(',', '.').replace(' ', ''))
            return v
        except (ValueError, TypeError):
            return None

    def set_valeur(col_id, valeur):
        cell = cellules.get(col_id)
        if not cell:
            cell = AssuranceCellule(ligne_id=ligne.id, colonne_id=col_id, valeur=valeur)
            cellules[col_id] = cell
            db.session.add(cell)
        else:
            cell.valeur = valeur

    for col in colonnes:
        if not col.formule:
            continue
        params = _json.loads(col.formule_params) if col.formule_params else {}
        if col.formule == 'prix_m2':
            cotisation = valeur_num(params.get('cotisation'))
            superficie = valeur_num(params.get('superficie'))
            if cotisation is not None and superficie and superficie != 0:
                set_valeur(col.id, round(cotisation / superficie, 4))
            else:
                set_valeur(col.id, '')
        elif col.formule == 'evolution_n':
            annee = valeur_num(params.get('annee'))
            ref = valeur_num(params.get('ref'))
            if annee is not None and ref is not None and ref != 0:
                set_valeur(col.id, round((annee - ref) / ref, 4))
            else:
                set_valeur(col.id, '')
        elif col.formule == 'evolution_base':
            annee = valeur_num(params.get('annee'))
            base = valeur_num(params.get('base'))
            if annee is not None and base is not None and base != 0:
                set_valeur(col.id, round((annee - base) / base, 4))
            else:
                set_valeur(col.id, '')


@app.route('/contrats/assurance/ligne/add', methods=['POST'])
def assurance_add_ligne():
    """Ajoute une nouvelle ligne (copropriété-assurance)."""
    max_ordre = db.session.query(db.func.max(AssuranceLigne.ordre)).scalar() or 0
    # On insère avant la ligne TOTAL si elle existe
    total = AssuranceLigne.query.filter_by(est_total=True).first()
    if total:
        # décale la TOTAL à la fin
        total.ordre = max_ordre + 2
        new_ordre = max_ordre + 1
    else:
        new_ordre = max_ordre + 1
    ligne = AssuranceLigne(ordre=new_ordre, est_total=False)
    db.session.add(ligne)
    db.session.flush()
    db.session.commit()
    flash('Nouvelle ligne ajoutée. Cliquez sur les cellules pour la remplir.', 'success')
    return redirect(url_for('assurance_page'))


@app.route('/contrats/assurance/ligne/<int:ligne_id>/delete', methods=['POST'])
def assurance_delete_ligne(ligne_id):
    ligne = AssuranceLigne.query.get_or_404(ligne_id)
    db.session.delete(ligne)
    db.session.commit()
    flash('Ligne supprimée.', 'success')
    return redirect(url_for('assurance_page'))


@app.route('/contrats/assurance/colonne/add', methods=['POST'])
def assurance_add_colonne():
    """Ajoute une nouvelle colonne (année/indicateur)."""
    groupe = request.form.get('groupe') or 'Tarif'
    en_tete = request.form.get('en_tete', '').strip()
    if not en_tete:
        flash('Veuillez donner un nom à la nouvelle colonne.', 'error')
        return redirect(url_for('assurance_page'))
    max_ordre = db.session.query(db.func.max(AssuranceColonne.ordre)).scalar() or 0
    col = AssuranceColonne(groupe=groupe, en_tete=en_tete, ordre=max_ordre + 1)
    db.session.add(col)
    db.session.commit()
    flash(f'Colonne « {en_tete} » ajoutée.', 'success')
    return redirect(url_for('assurance_page'))


@app.route('/contrats/assurance/colonne/<int:colonne_id>/delete', methods=['POST'])
def assurance_delete_colonne(colonne_id):
    col = AssuranceColonne.query.get_or_404(colonne_id)
    db.session.delete(col)
    db.session.commit()
    flash('Colonne supprimée.', 'success')
    return redirect(url_for('assurance_page'))


@app.route('/contrats/assurance/cellule/save', methods=['POST'])
def assurance_save_cellule():
    """Sauvegarde la valeur d'une cellule (édition inline)."""
    ligne_id = request.form.get('ligne_id', type=int)
    colonne_id = request.form.get('colonne_id', type=int)
    valeur = request.form.get('valeur', '')
    if ligne_id is None or colonne_id is None:
        flash('Données invalides pour la cellule.', 'error')
        return redirect(url_for('assurance_page'))
    cellule = AssuranceCellule.query.filter_by(
        ligne_id=ligne_id, colonne_id=colonne_id).first()
    if cellule:
        cellule.valeur = valeur
    else:
        cellule = AssuranceCellule(ligne_id=ligne_id, colonne_id=colonne_id, valeur=valeur)
        db.session.add(cellule)
    # Recalcule les colonnes formules de cette ligne si la cellule modifiée
    # est une colonne saisie (cotisation / superficie / base).
    ligne = AssuranceLigne.query.get(ligne_id)
    if ligne and not ligne.est_total:
        _recalculer_ligne(ligne)
    db.session.commit()
    flash('Cellule sauvegardée et recalculée.', 'success')
    return redirect(url_for('assurance_page'))

# ========== ONGLET HONORAIRES (tableau « Honoraire agate ») ==========

def _fmt_euro(valeur):
    """Formate un montant en euros avec le signe € (ex: 1 296,00 €)."""
    if valeur is None:
        return ''
    try:
        nombre = float(valeur)
    except (ValueError, TypeError):
        return str(valeur)
    texte = f"{nombre:,.2f}".replace(",", " ").replace(".", ",")
    return f"{texte} €"


@app.template_filter('euro')
def _euro_filter(valeur):
    return _fmt_euro(valeur)


@app.template_filter('fin_contrat_courte')
def _fin_contrat_courte_filter(valeur):
    """Affiche une date de fin de contrat au format court « 30 juin » / « 31 déc. »."""
    MOIS_COURTS = {
        1: 'janv.', 2: 'févr.', 3: 'mars', 4: 'avr.', 5: 'mai', 6: 'juin',
        7: 'juil.', 8: 'août', 9: 'sept.', 10: 'oct.', 11: 'nov.', 12: 'déc.',
    }
    if not valeur:
        return ''
    return f"{valeur.day} {MOIS_COURTS[valeur.month]}"


def _importer_honoraires():
    """Importe les lignes du tableau « Honoraire agate » (une seule fois,
    si la table est vide). Une fois fait, le fichier Excel peut être supprimé."""
    if HonoraireLigne.query.count() > 0:
        return
    for i, d in enumerate(HONORAIRE_DATA):
        ligne = HonoraireLigne(
            numero=d['numero'],
            ordre=i,
            mise_copro=parse_date(d['mise_copro']) if d['mise_copro'] else None,
            fin_contrat=parse_date(d['fin_contrat']) if d['fin_contrat'] else None,
            type_facturation=d['type'],
            logements=d['logements'],
            tarif_s1=d['tarif_s1'],
            tarif_s2=d['tarif_s2'],
            t12=d['t12'],
            t34=d['t34'],
            total_ttc=d['total_ttc'],
            total_ht=d['total_ht'],
        )
        db.session.add(ligne)
    db.session.commit()


def _recalculer_ligne_honoraire(ligne, tarif_s1=None, tarif_s2=None, t12=None, t34=None):
    """Recalcule les colonnes t12 / t34 / total_ttc / total_ht d'une ligne.

    Lignes Classique / Forfait / ASL : t12 et t34 sont calculés à partir
    des tarifs semestriels et du nombre de logements.

    Lignes « Spécifique » : t12 et t34 sont saisis à la main ; seuls les
    totaux TTC et HT sont recalculés à partir de ces montants.
    """
    if tarif_s1 is not None:
        ligne.tarif_s1 = tarif_s1
    if tarif_s2 is not None:
        ligne.tarif_s2 = tarif_s2
    if ligne.est_specifique:
        if t12 is not None:
            ligne.t12 = t12
        if t34 is not None:
            ligne.t34 = t34
    else:
        logements = ligne.logements or 0
        if ligne.tarif_s1 is not None:
            ligne.t12 = round(ligne.tarif_s1 * logements / 4, 2)
        if ligne.tarif_s2 is not None:
            ligne.t34 = round(ligne.tarif_s2 * logements / 4, 2)
    if ligne.t12 is not None or ligne.t34 is not None:
        ligne.total_ttc = round((ligne.t12 or 0) * 2 + (ligne.t34 or 0) * 2, 2)
        ligne.total_ht = round(ligne.total_ttc / 1.2, 2)
    else:
        ligne.total_ttc = None
        ligne.total_ht = None


@app.route('/honoraires', endpoint='honoraires_page')
def honoraires_page():
    """Tableau des honoraires (« Honoraire agate ») chargé depuis la base."""
    _importer_honoraires()
    lignes = HonoraireLigne.query.filter_by(est_total=False).order_by(
        HonoraireLigne.numero, HonoraireLigne.id).all()
    total_t12 = sum(l.t12 or 0 for l in lignes)
    total_t34 = sum(l.t34 or 0 for l in lignes)
    total_ttc = sum(l.total_ttc or 0 for l in lignes)
    total_ht = sum(l.total_ht or 0 for l in lignes)
    return render_template(
        'honoraires.html',
        lignes=lignes,
        types=TYPE_FACTURATION,
        total_t12_str=_fmt_euro(total_t12),
        total_t34_str=_fmt_euro(total_t34),
        total_ttc_str=_fmt_euro(total_ttc),
        total_ht_str=_fmt_euro(total_ht),
    )


@app.route('/honoraires/ligne/add', methods=['POST'], endpoint='honoraires_add_ligne')
def honoraires_add_ligne():
    """Ajoute une ligne d'honoraires et recalcule ses montants."""
    def f(name):
        v = request.form.get(name, '').strip()
        return v if v else None
    numero_raw = f('numero')
    numero = None
    if numero_raw is not None:
        try:
            numero = int(float(numero_raw.replace(',', '.')))
        except (ValueError, TypeError):
            numero = None
    if numero is None:
        flash('Veuillez indiquer un N° de copro valide.', 'error')
        return redirect(url_for('honoraires_page'))
    logements = None
    if f('logements'):
        try:
            logements = int(float(f('logements').replace(',', '.')))
        except (ValueError, TypeError):
            logements = None

    def money(name):
        v = f(name)
        if v is None:
            return None
        try:
            return float(v.replace('€', '').replace(' ', '').replace(',', '.'))
        except (ValueError, TypeError):
            return None

    ligne = HonoraireLigne(
        numero=numero,
        mise_copro=parse_date(f('mise_copro')),
        fin_contrat=parse_date(f('fin_contrat')),
        type_facturation=f('type_facturation'),
        logements=logements,
    )
    db.session.add(ligne)
    _recalculer_ligne_honoraire(
        ligne,
        tarif_s1=money('tarif_s1'),
        tarif_s2=money('tarif_s2'),
        t12=money('t12'),
        t34=money('t34'),
    )
    db.session.commit()
    flash(f"Ligne d'honoraires ajoutée pour la copropriété {numero}.", 'success')
    return redirect(url_for('honoraires_page'))


@app.route('/honoraires/ligne/<int:ligne_id>/delete', methods=['POST'], endpoint='honoraires_delete_ligne')
def honoraires_delete_ligne(ligne_id):
    ligne = HonoraireLigne.query.get_or_404(ligne_id)
    db.session.delete(ligne)
    db.session.commit()
    flash("Ligne d'honoraires supprimée.", 'success')
    return redirect(url_for('honoraires_page'))


@app.route('/honoraires/ligne/save', methods=['POST'], endpoint='honoraires_save_ligne')
def honoraires_save_ligne():
    """Sauvegarde une ligne complète et recalcule les colonnes calculées."""
    ligne_id = request.form.get('ligne_id', type=int)
    if ligne_id is None:
        flash('Ligne introuvable.', 'error')
        return redirect(url_for('honoraires_page'))
    ligne = HonoraireLigne.query.get_or_404(ligne_id)

    def money(name, default=None):
        v = request.form.get(name, '').strip()
        if not v:
            return default
        try:
            return float(v.replace('€', '').replace(' ', '').replace(',', '.'))
        except (ValueError, TypeError):
            return default

    ligne.mise_copro = parse_date(request.form.get('mise_copro', '').strip()) or ligne.mise_copro
    ligne.fin_contrat = parse_date(request.form.get('fin_contrat', '').strip()) or ligne.fin_contrat
    type_facturation = request.form.get('type_facturation', '').strip()
    if type_facturation:
        ligne.type_facturation = type_facturation
    if request.form.get('logements', '').strip():
        try:
            ligne.logements = int(float(request.form.get('logements', '').strip().replace(',', '.')))
        except (ValueError, TypeError):
            pass
    _recalculer_ligne_honoraire(
        ligne,
        tarif_s1=money('tarif_s1', ligne.tarif_s1),
        tarif_s2=money('tarif_s2', ligne.tarif_s2),
        t12=money('t12', ligne.t12),
        t34=money('t34', ligne.t34),
    )
    db.session.commit()
    flash('Honoraires sauvegardés et recalculés.', 'success')
    return redirect(url_for('honoraires_page'))


@app.route('/honoraires/export', endpoint='honoraires_export')
def honoraires_export():
    """Exporte le tableau des honoraires en Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    lignes = HonoraireLigne.query.filter_by(est_total=False).order_by(
        HonoraireLigne.numero, HonoraireLigne.id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Honoraires"

    headers = [
        'N° de copro', 'Mise en copro', 'Fin de contrat', 'Type de facturation',
        'Logements', '€ TTC / LP\n1er semestre', '€ TTC / LP\n2e semestre',
        '€ TTC\n1er et 2e\nTrimestre', '€ TTC\n3e et 4e\nTrimestre',
        'Total € TTC', 'Total € HT',
    ]
    gras_blanc = Font(bold=True, color="FFFFFF")
    fond = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = gras_blanc
        cell.fill = fond
        cell.alignment = centre
        cell.border = bordure

    ligne_excel = 2
    for l in lignes:
        valeurs = [
            l.numero,
            l.mise_copro.strftime('%d/%m/%Y') if l.mise_copro else '',
            l.fin_contrat.strftime('%d/%m/%Y') if l.fin_contrat else '',
            l.type_facturation or '',
            l.logements,
            l.tarif_s1, l.tarif_s2, l.t12, l.t34, l.total_ttc, l.total_ht,
        ]
        for c, v in enumerate(valeurs, start=1):
            cell = ws.cell(row=ligne_excel, column=c, value=v)
            cell.border = bordure
        ligne_excel += 1

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    ws.freeze_panes = "B2"

    nom_fichier = f"honoraires_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return send_file(
        tampon,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _honoraire_pour_copropriete(numero):
    """Récupère la ligne d'honoraires correspondant au N° de copro."""
    return HonoraireLigne.query.filter(
        HonoraireLigne.numero == numero, HonoraireLigne.est_total == False
    ).order_by(HonoraireLigne.id).first()


# ========== ONGLET VISITES D'IMMEUBLE (tableau « Visites d'immeuble agate ») ==========

def _normaliser_visite(valeur):
    """Convertit une date ISO (YYYY-MM-DD) au format JJ/MM/AAAA ; laisse le
    texte libre (ex. « A voir ») inchangé."""
    valeur = (valeur or '').strip()
    try:
        return datetime.strptime(valeur, '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return valeur


def _migrer_coproprietaires():
    """Ajoute la colonne est_conseil_syndical aux bases créées avant son
    introduction (db.create_all ne migre pas les tables existantes)."""
    with db.engine.begin() as conn:
        colonnes = [c['name'] for c in db.inspect(conn).get_columns('coproprietaire')]
        if 'est_conseil_syndical' not in colonnes:
            conn.execute(db.text('ALTER TABLE coproprietaire ADD COLUMN est_conseil_syndical BOOLEAN DEFAULT 0'))


def _migrer_visite_lignes():
    """Ajoute la colonne mise_copro aux bases créées avant son introduction
    (db.create_all ne migre pas les tables existantes)."""
    with db.engine.begin() as conn:
        colonnes = [c['name'] for c in db.inspect(conn).get_columns('visite_lignes')]
        if 'mise_copro' not in colonnes:
            conn.execute(db.text('ALTER TABLE visite_lignes ADD COLUMN mise_copro VARCHAR(20)'))


def _migrer_travaux_appels():
    """Ajoute la colonne fait aux appels de fonds des bases créées avant son
    introduction (db.create_all ne migre pas les tables existantes)."""
    with db.engine.begin() as conn:
        colonnes = [c['name'] for c in db.inspect(conn).get_columns('travaux_appels')]
        if 'fait' not in colonnes:
            conn.execute(db.text('ALTER TABLE travaux_appels ADD COLUMN fait BOOLEAN DEFAULT 0'))


def _completer_mise_copro():
    """Complète les dates de mise en copro manquantes depuis les données
    importées de l'Excel."""
    par_numero = {l['numero']: l.get('mise_copro', '') for l in VISITES_LIGNES}
    maj = False
    for ligne in VisiteLigne.query.all():
        if (ligne.mise_copro or '').strip():
            continue
        valeur = _normaliser_visite(par_numero.get(ligne.numero, ''))
        if valeur:
            ligne.mise_copro = valeur
            maj = True
    if maj:
        db.session.commit()


def _importer_visites():
    """Importe les lignes du tableau « Visites d'immeuble agate » (une seule
    fois, si les tables sont vides). Une fois fait, le fichier Excel peut
    être supprimé."""
    _migrer_coproprietaires()
    _migrer_visite_lignes()
    if VisiteColonne.query.count() > 0 or VisiteLigne.query.count() > 0:
        _completer_mise_copro()
        return
    colonnes = {}
    for i, annee in enumerate(VISITES_ANNEES):
        col = VisiteColonne(annee=annee, ordre=i)
        db.session.add(col)
        colonnes[annee] = col
    db.session.flush()
    for i, l in enumerate(VISITES_LIGNES):
        ligne = VisiteLigne(numero=l['numero'], ordre=i,
                            mise_copro=_normaliser_visite(l.get('mise_copro', '')))
        db.session.add(ligne)
        db.session.flush()
        for annee, valeur in l['visites'].items():
            col = colonnes.get(annee)
            if col is None:
                continue
            db.session.add(VisiteCellule(
                ligne_id=ligne.id, colonne_id=col.id,
                valeur=_normaliser_visite(valeur)))
    db.session.commit()


def _visites_context():
    """Construit le contexte de la page Visites depuis la base."""
    colonnes = VisiteColonne.query.order_by(VisiteColonne.ordre).all()
    lignes = VisiteLigne.query.order_by(VisiteLigne.numero, VisiteLigne.id).all()
    rows = []
    for ligne in lignes:
        cellules = {c.colonne_id: c.valeur for c in ligne.cellules}
        rows.append({
            'id': ligne.id,
            'numero': ligne.numero,
            'mise_copro': ligne.mise_copro or '',
            'valeurs': [cellules.get(col.id, '') for col in colonnes],
        })
    return {'colonnes': colonnes, 'rows': rows}


def _visites_pour_copropriete(numero):
    """Renvoie les visites d'une copropriété : liste de dicts {colonne_id,
    annee, valeur}. Ne renvoie QUE les cellules non vides."""
    colonnes = VisiteColonne.query.order_by(VisiteColonne.ordre).all()
    ligne = VisiteLigne.query.filter(VisiteLigne.numero == numero).order_by(
        VisiteLigne.id).first()
    if not ligne:
        return None
    cellules = {c.colonne_id: c.valeur for c in ligne.cellules}
    visites = []
    for col in colonnes:
        valeur = (cellules.get(col.id) or '').strip()
        if valeur:
            visites.append({'colonne_id': col.id, 'annee': col.annee, 'valeur': valeur})
    return {'ligne_id': ligne.id, 'visites': visites}


@app.route('/visites', endpoint='visites_page')
def visites_page():
    """Tableau des visites d'immeuble chargé depuis la base."""
    _importer_visites()
    return render_template('visites.html', **_visites_context())


@app.route('/visites/ligne/add', methods=['POST'], endpoint='visites_add_ligne')
def visites_add_ligne():
    """Ajoute une ligne (copropriété) au tableau des visites."""
    numero_raw = request.form.get('numero', '').strip()
    try:
        numero = int(float(numero_raw.replace(',', '.')))
    except (ValueError, TypeError):
        flash('Veuillez indiquer un N° de copro valide.', 'error')
        return redirect(url_for('visites_page'))
    max_ordre = db.session.query(db.func.max(VisiteLigne.ordre)).scalar() or 0
    ligne = VisiteLigne(numero=numero, ordre=max_ordre + 1)
    db.session.add(ligne)
    db.session.commit()
    flash(f'Ligne ajoutée pour la copropriété {numero}.', 'success')
    return redirect(url_for('visites_page'))


@app.route('/visites/ligne/<int:ligne_id>/delete', methods=['POST'], endpoint='visites_delete_ligne')
def visites_delete_ligne(ligne_id):
    ligne = VisiteLigne.query.get_or_404(ligne_id)
    db.session.delete(ligne)
    db.session.commit()
    flash('Ligne supprimée.', 'success')
    return redirect(url_for('visites_page'))


@app.route('/visites/ligne/<int:ligne_id>/mise-copro', methods=['POST'], endpoint='visites_save_mise_copro')
def visites_save_mise_copro(ligne_id):
    """Sauvegarde la date de mise en copro d'une ligne (JJ/MM/AAAA)."""
    ligne = VisiteLigne.query.get_or_404(ligne_id)
    ligne.mise_copro = (request.form.get('valeur') or '').strip()
    db.session.commit()
    flash('Date de mise en copro sauvegardée.', 'success')
    return redirect(request.referrer or url_for('visites_page'))


@app.route('/visites/colonne/add', methods=['POST'], endpoint='visites_add_colonne')
def visites_add_colonne():
    """Ajoute une colonne (année de visite)."""
    annee = request.form.get('annee', '').strip()
    if not annee:
        flash("Veuillez indiquer l'année de la nouvelle colonne.", 'error')
        return redirect(url_for('visites_page'))
    if VisiteColonne.query.filter_by(annee=annee).first():
        flash(f'La colonne {annee} existe déjà.', 'error')
        return redirect(url_for('visites_page'))
    max_ordre = db.session.query(db.func.max(VisiteColonne.ordre)).scalar() or 0
    col = VisiteColonne(annee=annee, ordre=max_ordre + 1)
    db.session.add(col)
    db.session.commit()
    flash(f'Colonne {annee} ajoutée.', 'success')
    return redirect(url_for('visites_page'))


@app.route('/visites/colonne/<int:colonne_id>/delete', methods=['POST'], endpoint='visites_delete_colonne')
def visites_delete_colonne(colonne_id):
    col = VisiteColonne.query.get_or_404(colonne_id)
    db.session.delete(col)
    db.session.commit()
    flash('Colonne supprimée.', 'success')
    return redirect(url_for('visites_page'))


def _visite_cellule_get(ligne_id, colonne_id):
    return VisiteCellule.query.filter_by(
        ligne_id=ligne_id, colonne_id=colonne_id).first()


@app.route('/visites/cellule/save', methods=['POST'], endpoint='visites_save_cellule')
def visites_save_cellule():
    """Sauvegarde une cellule (date de visite ou texte libre).
    Valeur vide = suppression de la cellule (aucune visite)."""
    ligne_id = request.form.get('ligne_id', type=int)
    colonne_id = request.form.get('colonne_id', type=int)
    valeur = (request.form.get('valeur') or '').strip()
    if ligne_id is None or colonne_id is None:
        flash('Données invalides pour la cellule.', 'error')
        return redirect(request.referrer or url_for('visites_page'))
    cellule = _visite_cellule_get(ligne_id, colonne_id)
    if valeur:
        if cellule:
            cellule.valeur = valeur
        else:
            db.session.add(VisiteCellule(
                ligne_id=ligne_id, colonne_id=colonne_id, valeur=valeur))
    elif cellule:
        db.session.delete(cellule)
    db.session.commit()
    flash('Visite sauvegardée.', 'success')
    dest = request.form.get('from') or 'visites'
    if dest == 'copropriete':
        copro_id = request.form.get('copro_id', type=int)
        if copro_id:
            return redirect(url_for('copropriete', copro_id=copro_id))
    return redirect(url_for('visites_page'))


@app.route('/visites/export', endpoint='visites_export')
def visites_export():
    """Exporte le tableau des visites en Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    colonnes = VisiteColonne.query.order_by(VisiteColonne.ordre).all()
    lignes = VisiteLigne.query.order_by(VisiteLigne.numero, VisiteLigne.id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visite d'immeuble"

    gras_blanc = Font(bold=True, color="FFFFFF")
    fond = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ['N° de copro', 'Mise en copro'] + [c.annee for c in colonnes]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = gras_blanc
        cell.fill = fond
        cell.alignment = centre
        cell.border = bordure

    ligne_excel = 2
    for ligne in lignes:
        cellules = {cl.colonne_id: cl.valeur for cl in ligne.cellules}
        ws.cell(row=ligne_excel, column=1, value=ligne.numero).border = bordure
        ws.cell(row=ligne_excel, column=2, value=ligne.mise_copro or '').border = bordure
        for c, col in enumerate(colonnes, start=3):
            ws.cell(row=ligne_excel, column=c, value=cellules.get(col.id, '')).border = bordure
        ligne_excel += 1

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 16
    for c in range(3, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    ws.freeze_panes = "C2"

    nom_fichier = f"visites_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return send_file(
        tampon,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ========== ONGLET TRAVAUX (tableau « ADF Travaux AGATE ») ==========

def _date_fr_vers_iso(valeur):
    """Convertit JJ/MM/AAAA (ou JJ/MM/AA) en date ; renvoie None sinon.
    Accepte aussi les séparateurs « - » et « . » à la place des « / »."""
    texte = (valeur or '').strip().replace('-', '/').replace('.', '/')
    for fmt in ('%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(texte, fmt).date()
        except ValueError:
            continue
    return None


def _premiere_date_appel(ligne):
    """Date d'appel de fonds la plus ancienne d'une ligne (date, ou None)."""
    dates = []
    for a in ligne.appels:
        d = _date_fr_vers_iso(a.date_appel)
        if d:
            dates.append(d)
    return min(dates) if dates else None


def _devis_a_faire(ligne):
    """True si la case Devis validé est vide et que la date du jour a
    dépassé la date d'appel la plus ancienne de la ligne."""
    if (ligne.devis_valide or '').strip():
        return False
    d = _premiere_date_appel(ligne)
    return d is not None and date.today() > d


def _fin_exercice_courte(valeur_iso):
    """Convertit une date ISO (YYYY-MM-DD) en « 30 juin » / « 31 décembre »."""
    if not valeur_iso:
        return None
    d = datetime.strptime(valeur_iso, '%Y-%m-%d').date()
    if d.day == 30 and d.month == 6:
        return '30 juin'
    if d.day == 31 and d.month == 12:
        return '31 décembre'
    return d.strftime('%d/%m/%Y')


def _importer_travaux():
    """Importe les lignes du tableau « ADF Travaux AGATE » (une seule fois,
    si la table est vide). Une fois fait, le fichier Excel peut être supprimé."""
    if TravauxLigne.query.count() > 0:
        return
    for l in TRAVAUX_LIGNES:
        ligne = TravauxLigne(
            numero=l['numero'],
            fin_exercice=_fin_exercice_courte(l.get('fin_exercice')),
            date_ag=(datetime.strptime(l['date_ag'], '%Y-%m-%d').date()
                     if l.get('date_ag') else None),
            type=l.get('type'),
            montant=l.get('montant'),
            honoraire=l.get('honoraire'),
            nb_appels=l.get('nb_appels'),
            tethrawin=l.get('tethrawin'),
            devis_valide=l.get('devis_valide'),
            os=l.get('os'),
            facture=l.get('facture'),
            annee_cloture=l.get('annee_cloture'),
            statut=l.get('statut') or 'en_cours',
        )
        db.session.add(ligne)
    db.session.commit()


def _travaux_context():
    """Toutes les lignes de travaux triées par Date d'AG décroissante."""
    lignes = TravauxLigne.query.order_by(TravauxLigne.date_ag.desc().nullslast(),
                                         TravauxLigne.id).all()
    for ligne in lignes:
        ligne.devis_a_faire = _devis_a_faire(ligne)
        ligne.date_ag_str = ligne.date_ag.strftime('%d/%m/%Y') if ligne.date_ag else ''
        ligne.fin_exercice_str = ligne.fin_exercice or ''
        ligne.premiere_date_appel = _premiere_date_appel(ligne)
    return {'lignes': lignes}


def _travaux_pour_copropriete(numero):
    """Lignes de travaux d'une copropriété (triées par Date d'AG)."""
    lignes = TravauxLigne.query.filter(TravauxLigne.numero == numero).order_by(
        TravauxLigne.date_ag.desc().nullslast(), TravauxLigne.id).all()
    for ligne in lignes:
        ligne.devis_a_faire = _devis_a_faire(ligne)
        ligne.date_ag_str = ligne.date_ag.strftime('%d/%m/%Y') if ligne.date_ag else ''
        ligne.fin_exercice_str = ligne.fin_exercice or ''
    return lignes


@app.route('/travaux', endpoint='travaux_page')
def travaux_page():
    """Tableau des travaux + tableaux de suivi (devis à valider, appels
    de fonds à faire pour la comptable)."""
    _importer_travaux()
    _migrer_travaux_appels()
    ctx = _travaux_context()
    aujourdhui = date.today()
    limite = aujourdhui - timedelta(days=62)
    suivi_devis = [l for l in ctx['lignes'] if l.devis_a_faire]
    # Suivi comptable : une ligne par date d'appel, qu'elle soit faite ou
    # non. La ligne ne disparaît qu'une fois la date d'appel antérieure
    # de plus de deux mois à la date du jour.
    def _appel_a_suivre(ligne, appel):
        d = _date_fr_vers_iso(appel.date_appel)
        return d is not None and d >= limite

    suivi_appels = []
    for l in ctx['lignes']:
        for a in l.appels:
            if _appel_a_suivre(l, a):
                suivi_appels.append((l, a))
    suivi_appels.sort(key=lambda pa: _date_fr_vers_iso(pa[1].date_appel) or date.max)
    return render_template('travaux.html', **ctx,
                           suivi_devis=suivi_devis, suivi_appels=suivi_appels)


@app.route('/travaux/ligne/add', methods=['POST'], endpoint='travaux_add_ligne')
def travaux_add_ligne():
    """Ajoute une ligne de travaux."""
    numero_raw = request.form.get('numero', '').strip()
    try:
        numero = int(float(numero_raw.replace(',', '.')))
    except (ValueError, TypeError):
        flash('Veuillez indiquer un N° de copro valide.', 'error')
        return redirect(url_for('travaux_page'))
    ligne = TravauxLigne(numero=numero, statut='en_cours')
    ligne.date_ag = parse_date(request.form.get('date_ag'))
    ligne.type = request.form.get('type', '').strip() or None
    montant = request.form.get('montant', '').replace('€', '').replace(' ', '').replace(',', '.')
    if montant:
        try:
            ligne.montant = float(montant)
        except ValueError:
            pass
    db.session.add(ligne)
    db.session.commit()
    flash(f'Ligne de travaux ajoutée pour la copropriété {numero}.', 'success')
    if request.form.get('from') == 'copropriete':
        copro_id = request.form.get('copro_id', type=int)
        if copro_id:
            return redirect(url_for('copropriete', copro_id=copro_id))
    return redirect(url_for('travaux_page'))


@app.route('/travaux/ligne/<int:ligne_id>/delete', methods=['POST'], endpoint='travaux_delete_ligne')
def travaux_delete_ligne(ligne_id):
    ligne = TravauxLigne.query.get_or_404(ligne_id)
    db.session.delete(ligne)
    db.session.commit()
    flash('Ligne de travaux supprimée.', 'success')
    return redirect(url_for('travaux_page'))


@app.route('/travaux/ligne/<int:ligne_id>/save', methods=['POST'], endpoint='travaux_save_ligne')
def travaux_save_ligne(ligne_id):
    """Sauvegarde d'une cellule ou d'une ligne entière depuis le tableau
    ou la fiche copropriété."""
    ligne = TravauxLigne.query.get_or_404(ligne_id)
    champ = request.form.get('champ')
    if champ:
        valeur = (request.form.get('valeur') or '').strip()
        if champ == 'date_ag':
            ligne.date_ag = parse_date(valeur) if valeur else None
        elif champ == 'fin_exercice':
            ligne.fin_exercice = valeur or None
        elif champ in ('montant', 'honoraire'):
            if valeur:
                try:
                    setattr(ligne, champ, float(valeur.replace('€', '').replace(' ', '').replace(',', '.')))
                except ValueError:
                    pass
            else:
                setattr(ligne, champ, None)
        elif champ == 'nb_appels':
            if valeur:
                try:
                    ligne.nb_appels = int(valeur)
                except ValueError:
                    pass
            else:
                ligne.nb_appels = None
        elif champ == 'statut':
            ligne.statut = valeur if valeur in STATUTS_TRAVAUX else 'en_cours'
        elif champ == 'honoraire_facture':
            ligne.honoraire_facture = valeur == '1'
        elif champ == 'fait':
            ligne.fait = valeur == '1'
        elif hasattr(ligne, champ) and champ not in ('id', 'numero'):
            setattr(ligne, champ, valeur or None)
        db.session.commit()
        flash('Travaux sauvegardé.', 'success')
    if request.form.get('from') == 'copropriete':
        copro_id = request.form.get('copro_id', type=int)
        if copro_id:
            return redirect(url_for('copropriete', copro_id=copro_id))
    return redirect(url_for('travaux_page'))


@app.route('/travaux/appel/save', methods=['POST'], endpoint='travaux_save_appel')
def travaux_save_appel():
    """Sauvegarde une date d'appel de fonds. Empty valeur = suppression."""
    ligne_id = request.form.get('ligne_id', type=int)
    appel_id = request.form.get('appel_id', type=int)
    valeur = (request.form.get('valeur') or '').strip()
    ligne = TravauxLigne.query.get_or_404(ligne_id)
    if appel_id:
        appel = TravauxAppel.query.get_or_404(appel_id)
        if valeur:
            appel.date_appel = valeur
        else:
            db.session.delete(appel)
    elif valeur:
        max_ordre = db.session.query(db.func.max(TravauxAppel.ordre)).filter_by(
            ligne_id=ligne_id).scalar() or 0
        db.session.add(TravauxAppel(ligne_id=ligne_id, ordre=max_ordre + 1,
                                    date_appel=valeur))
    db.session.commit()
    if request.form.get('from') == 'copropriete':
        copro_id = request.form.get('copro_id', type=int)
        if copro_id:
            return redirect(url_for('copropriete', copro_id=copro_id))
    return redirect(url_for('travaux_page'))


@app.route('/travaux/appel/<int:appel_id>/fait', methods=['POST'], endpoint='travaux_appel_fait')
def travaux_appel_fait(appel_id):
    """Coche/décoche « Fait » pour une date d'appel de fonds du suivi comptable."""
    _migrer_travaux_appels()
    appel = TravauxAppel.query.get_or_404(appel_id)
    appel.fait = not appel.fait
    db.session.commit()
    return redirect(url_for('travaux_page'))


@app.route('/travaux/ligne/<int:ligne_id>/appels', methods=['POST'], endpoint='travaux_save_appels')
def travaux_save_appels(ligne_id):
    """Remplace toutes les dates d'appels de fonds d'une ligne par celles
    soumises (champs date_0, date_1, ...). Une date vide est ignorée.
    Le champ optionnel « nb » met à jour le nombre d'appels."""
    ligne = TravauxLigne.query.get_or_404(ligne_id)
    nb = request.form.get('nb', '').strip()
    dates = []
    i = 0
    while f'date_{i}' in request.form:
        valeur = (request.form.get(f'date_{i}') or '').strip()
        if valeur:
            dates.append(valeur)
        i += 1
    faits_existants = {a.date_appel: a.fait for a in TravauxAppel.query.filter_by(ligne_id=ligne_id).all()}
    TravauxAppel.query.filter_by(ligne_id=ligne_id).delete()
    for ordre, d in enumerate(dates):
        db.session.add(TravauxAppel(ligne_id=ligne_id, ordre=ordre, date_appel=d,
                                    fait=faits_existants.get(d, False)))
    if nb:
        try:
            ligne.nb_appels = int(nb)
        except ValueError:
            pass
    elif not ligne.nb_appels or ligne.nb_appels < len(dates):
        ligne.nb_appels = len(dates)
    db.session.commit()
    flash('Dates des appels de fonds sauvegardées.', 'success')
    if request.form.get('from') == 'copropriete':
        copro_id = request.form.get('copro_id', type=int)
        if copro_id:
            return redirect(url_for('copropriete', copro_id=copro_id))
    return redirect(url_for('travaux_page'))


@app.route('/travaux/export', endpoint='travaux_export')
def travaux_export():
    """Exporte le tableau des travaux en Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    lignes = _travaux_context()['lignes']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Travaux"

    gras_blanc = Font(bold=True, color="FFFFFF")
    fond = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        'N° de copro', 'Fin exercice', "Date d'AG", 'Type', 'Montant', 'Honoraire',
        'Honoraire facturé', "Nombre d'appels", 'Appels (dates)', 'Téthrawin',
        'Devis validé', 'OS', 'Facture', 'Année de clôture', 'Clôturé',
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = gras_blanc
        cell.fill = fond
        cell.alignment = centre
        cell.border = bordure

    r = 2
    for l in lignes:
        valeurs = [
            l.numero,
            l.fin_exercice or '',
            l.date_ag.strftime('%d/%m/%Y') if l.date_ag else '',
            l.type or '',
            l.montant if l.montant is not None else '',
            l.honoraire if l.honoraire is not None else '',
            'Oui' if l.honoraire_facture else '',
            l.nb_appels if l.nb_appels is not None else '',
            ', '.join(a.date_appel or '' for a in l.appels),
            l.tethrawin or '',
            ('A FAIRE' if l.devis_a_faire else (l.devis_valide or '')),
            l.os or '',
            l.facture or '',
            l.annee_cloture or '',
            {'en_cours': 'En cours', 'pret_ag': 'Prêt pour AG', 'termine': 'Terminé'}.get(l.statut, l.statut),
        ]
        for c, v in enumerate(valeurs, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = bordure
        r += 1

    largeurs = [12, 14, 12, 30, 12, 12, 14, 14, 24, 14, 14, 18, 18, 14, 14]
    for c, w in enumerate(largeurs, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    nom_fichier = f"travaux_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return send_file(
        tampon,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route('/')
def index():
    stats = get_statistiques()
    return render_template('index.html', stats=stats)


@app.route('/coproprietes', endpoint='coproprietes_liste')
def coproprietes_liste():
    search_query = request.args.get('q', '')
    search_type = request.args.get('type', 'tout')

    query = Copropriete.query.order_by(Copropriete.numero)

    if search_query:
        if search_type == 'tout':
            query = query.filter(
                (Copropriete.numero.ilike(f'%{search_query}%')) |
                (Copropriete.nom.ilike(f'%{search_query}%')) |
                (Copropriete.adresse.ilike(f'%{search_query}%')) |
                (Copropriete.ville.ilike(f'%{search_query}%')) |
                (Copropriete.gestionnaire.ilike(f'%{search_query}%')) |
                (Copropriete.immatriculation.ilike(f'%{search_query}%'))
            )
        elif search_type == 'ville':
            query = query.filter(Copropriete.ville.ilike(f'%{search_query}%'))
        elif search_type == 'gestionnaire':
            query = query.filter(Copropriete.gestionnaire.ilike(f'%{search_query}%'))
        elif search_type == 'comptable':
            query = query.filter(Copropriete.comptable.ilike(f'%{search_query}%'))
        elif search_type == 'immatriculation':
            query = query.filter(Copropriete.immatriculation.ilike(f'%{search_query}%'))

    coproprietes = query.all()
    return render_template('coproprietes_liste.html', coproprietes=coproprietes, search_query=search_query, search_type=search_type)


@app.route('/coproprietes/export', endpoint='coproprietes_export')
def coproprietes_export():
    """Exporte la liste des copropriétés en Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    coproprietes = Copropriete.query.order_by(Copropriete.numero).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Copropriétés"

    headers = [
        'N°', 'Nom', 'Adresse', 'Ville', 'Gestionnaire', 'Comptable',
        'Logements', 'Immatriculation', 'Statut',
    ]
    gras_blanc = Font(bold=True, color="FFFFFF")
    fond = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = gras_blanc
        cell.fill = fond
        cell.alignment = centre
        cell.border = bordure

    ligne_excel = 2
    for copro in coproprietes:
        valeurs = [
            copro.numero,
            copro.nom or '',
            copro.adresse or '',
            copro.ville or '',
            copro.gestionnaire or '',
            copro.comptable or '',
            copro.nombre_logements,
            copro.immatriculation or '',
            'Active' if copro.est_active else 'Inactive',
        ]
        for c_idx, v in enumerate(valeurs, start=1):
            cell = ws.cell(row=ligne_excel, column=c_idx, value=v)
            cell.border = bordure
        ligne_excel += 1

    largeurs = [6, 30, 35, 20, 14, 20, 11, 18, 11]
    for c_idx, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c_idx)].width = largeur
    ws.freeze_panes = "A2"

    nom_fichier = f"coproprietes_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return send_file(
        tampon,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

def _assurance_pour_copropriete(numero):
    """Récupère les données du contrat d'assurance d'une copropriété.

    La correspondance se fait sur la première colonne du tableau Assurance
    (« N° de copro ») qui contient le numéro de la copropriété.

    Renvoie un dict avec : entreprise, debut, numero_contrat, superficie,
    et le tarif + prix au m² de la dernière année tarifaire disponible
    (la plus à droite dont l'en-tête commence par « Tarif » et la cellule
    « Prix au m² » associée).
    Renvoie None si aucune ligne d'assurance ne correspond à ce numéro.
    """
    import re
    colonnes = AssuranceColonne.query.order_by(AssuranceColonne.ordre).all()
    if not colonnes:
        return None

    col_by_ordre = {i: c for i, c in enumerate(colonnes)}

    # Repère la ligne dont la première cellule vaut le numéro de la copro.
    cible = None
    for ligne in AssuranceLigne.query.order_by(AssuranceLigne.ordre).all():
        if ligne.est_total:
            continue
        cellules = {cl.colonne_id: cl.valeur for cl in ligne.cellules}
        premiere = cellules.get(colonnes[0].id, '')
        try:
            if premiere is not None and str(premiere).strip() == str(numero):
                cible = (ligne, cellules)
                break
        except (ValueError, TypeError):
            continue
    if cible is None:
        return None
    ligne, cellules = cible

    def val(ordre):
        col = col_by_ordre.get(ordre)
        if col is None:
            return ''
        return cellules.get(col.id, '') or ''

    entreprise = val(5)
    debut = val(6)
    numero_contrat = val(8)
    superficie = val(9)

    # Dernière colonne « Tarif AAAA » non vide = tarif de la dernière année.
    dernier_tarif_ordre = None
    dernier_tarif_annee = None
    for c in colonnes:
        m = re.match(r'^Tarif\s+(\d{4})$', c.en_tete.strip() if c.en_tete else '')
        if not m:
            continue
        v = cellules.get(c.id, '')
        if v not in (None, '', '0', 0):
            dernier_tarif_ordre = c.ordre
            dernier_tarif_annee = m.group(1)

    dernier_tarif = ''
    dernier_prix_m2 = ''
    if dernier_tarif_ordre is not None:
        dernier_tarif = val(dernier_tarif_ordre)
        # Colonne « Prix au m²/€TTC AAAA » située juste après le tarif de l'année.
        for c in colonnes:
            if (c.en_tete and c.formule == 'prix_m2'
                    and c.ordre > dernier_tarif_ordre
                    and dernier_tarif_annee and dernier_tarif_annee in c.en_tete):
                dernier_prix_m2 = cellules.get(c.id, '') or ''
                break

    return {
        'entreprise': entreprise,
        'debut': debut,
        'numero_contrat': numero_contrat,
        'superficie': superficie,
        'derniere_annee': dernier_tarif_annee,
        'dernier_tarif': dernier_tarif,
        'dernier_prix_m2': dernier_prix_m2,
    }


@app.route('/copropriete/<int:copro_id>')
def copropriete(copro_id):
    _migrer_coproprietaires()
    copropriete = Copropriete.query.options(
        joinedload(Copropriete.fiche_immeuble),
        joinedload(Copropriete.coproprietaires),
        joinedload(Copropriete.coproprietaires).joinedload(Coproprietaire.lots),
        joinedload(Copropriete.coproprietaires).joinedload(Coproprietaire.emails),
        joinedload(Copropriete.coproprietaires).joinedload(Coproprietaire.telephones),
        joinedload(Copropriete.contrats).joinedload(Contrat.prestations),
        joinedload(Copropriete.assemblees_generales).joinedload(AssembleeGenerale.points_a_retenir),
        joinedload(Copropriete.assemblees_generales).joinedload(AssembleeGenerale.budgets_travaux).joinedload(BudgetTravaux.appels_fonds),
        joinedload(Copropriete.resolutions_futures)
    ).get_or_404(copro_id)

    types_contrats = ["Assurance", "Nettoyage", "Entretien", "Sécurité", "Autre"]
    civilites = Civilite.query.all()
    assurance_contrat = _assurance_pour_copropriete(copropriete.numero)
    honoraire_ligne = _honoraire_pour_copropriete(copropriete.numero)
    visites = _visites_pour_copropriete(copropriete.numero)
    travaux = _travaux_pour_copropriete(copropriete.numero)

    return render_template(
        'copropriete.html',
        copropriete=copropriete,
        contrats=copropriete.contrats,
        coproprietaires=copropriete.coproprietaires,
        assemblees=copropriete.assemblees_generales,
        resolutions=copropriete.resolutions_futures,
        demandes=copropriete.demandes,
        types_contrats=types_contrats,
        civilites=civilites,
        assurance_contrat=assurance_contrat,
        honoraire_ligne=honoraire_ligne,
        visites=visites,
        travaux=travaux,
    )

@app.route('/coproprietaire/<int:coproprietaire_id>/delete', methods=['POST'])
def delete_coproprietaire(coproprietaire_id):
    coproprietaire = Coproprietaire.query.get_or_404(coproprietaire_id)
    copro_id = coproprietaire.copropriete_id
    db.session.delete(coproprietaire)
    db.session.commit()
    flash('Copropriétaire supprimé avec succès !', 'success')
    if request.form.get('from') == 'coproprietaires' or request.args.get('from') == 'coproprietaires':
        return redirect(url_for('coproprietaires_page'))
    return redirect(url_for('copropriete', copro_id=copro_id))

# ========== ONGLET COPROPRIÉTAIRES (vue globale) ==========
@app.route('/coproprietaires', endpoint='coproprietaires_page')
def coproprietaires_page():
    """Tableau dynamique de tous les copropriétaires, groupés par N° de copro."""
    _migrer_coproprietaires()
    search_query = request.args.get('q', '').strip()
    filter_copro = request.args.get('copro', '').strip()

    query = Coproprietaire.query.join(Copropriete)
    if filter_copro:
        try:
            num = int(filter_copro)
            query = query.filter(Copropriete.numero == num)
        except (ValueError, TypeError):
            query = query.filter(Copropriete.nom.ilike(f'%{filter_copro}%'))
    if search_query:
        like = f'%{search_query}%'
        query = query.filter(
            db.or_(
                Coproprietaire.nom.ilike(like),
                Coproprietaire.prenom.ilike(like),
                Coproprietaire.email.ilike(like),
                Coproprietaire.telephone.ilike(like),
                Coproprietaire.adresse.ilike(like),
                Coproprietaire.ville.ilike(like),
                Copropriete.nom.ilike(like),
            )
        )
    # Tri par N° de copro puis par id d'insertion : un copropriétaire ajouté
    # depuis la fiche copropriété apparaît après les copropriétaires existants
    # de la même copropriété (id croissant).
    query = query.order_by(Copropriete.numero, Coproprietaire.id)
    coproprietaires = query.all()

    copros = Copropriete.query.order_by(Copropriete.numero).all()
    civilites = Civilite.query.all()
    return render_template(
        'coproprietaires.html',
        coproprietaires=coproprietaires,
        copros=copros,
        civilites=civilites,
        search_query=search_query,
        filter_copro=filter_copro,
    )


@app.route('/coproprietaires/export', endpoint='coproprietaires_export')
def coproprietaires_export():
    """Exporte le tableau de l'onglet Copropriétaires en fichier Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    query = Coproprietaire.query.join(Copropriete).order_by(Copropriete.numero, Coproprietaire.id)
    coproprietaires = query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Copropriétaires"

    headers = [
        'N° de copro', 'Civilité', 'Nom', 'Prénom', 'Lots',
        'Adresse', 'Code postal', 'Ville',
        'Téléphones Principal', 'Téléphone Secondaire', 'Emails Principal',
    ]
    gras_blanc = Font(bold=True, color="FFFFFF")
    fond = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = gras_blanc
        cell.fill = fond
        cell.alignment = centre
        cell.border = bordure

    ligne = 2
    for cp in coproprietaires:
        lots_str = ', '.join(
            f"{lot.numero}{(' ' + lot.nature) if lot.nature else ''}" for lot in cp.lots
        ) or (cp.numero_lot or 'NC')
        email_principal = next((e.value for e in cp.emails if e.principal), None) or cp.email or ''
        email_secondaire = ', '.join(e.value for e in cp.emails if not e.principal)
        tel_principal = next((t.value for t in cp.telephones if t.principal), None) or cp.telephone or ''
        tel_secondaire = ', '.join(t.value for t in cp.telephones if not t.principal)
        valeurs = [
            cp.copropriete.numero,
            cp.civilite.libelle if cp.civilite else '',
            cp.nom or '',
            cp.prenom or '',
            lots_str,
            cp.adresse or '',
            cp.code_postal or '',
            cp.ville or '',
            tel_principal,
            tel_secondaire,
            email_principal,
        ]
        for c, v in enumerate(valeurs, start=1):
            cell = ws.cell(row=ligne, column=c, value=v)
            cell.border = bordure
        ligne += 1

    largeurs = [12, 20, 24, 24, 18, 30, 12, 18, 22, 22, 30]
    for c, w in enumerate(largeurs, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    nom_fichier = f"coproprietaires_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return send_file(
        tampon,
        as_attachment=True,
        download_name=nom_fichier,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route('/coproprietaire/<int:coproprietaire_id>/edit-inline', endpoint='coproprietaire_edit_inline')
def coproprietaire_edit_inline(coproprietaire_id):
    """Renvoie le contenu HTML d'une modale d'édition d'un copropriétaire,
    destiné à être chargé en AJAX depuis l'onglet Copropriétaires.
    """
    cp = Coproprietaire.query.get_or_404(coproprietaire_id)
    civilites = Civilite.query.all()
    return render_template('coproprietaire_edit_inline.html', cp=cp, civilites=civilites)


@app.route('/copropriete/new', methods=['GET', 'POST'])
def new_copropriete():
    if request.method == 'POST':
        numero = int(request.form.get('numero'))
        existing = Copropriete.query.filter_by(numero=numero).first()
        if existing:
            flash(f'Le numéro {numero} existe déjà !', 'error')
            return redirect(url_for('new_copropriete'))

        new_copro = Copropriete(
            numero=numero,
            nom=request.form.get('nom', f"Copropriété {numero}"),
            date_mise_copro=parse_date(request.form.get('date_mise_copropriete')),
            programme_neolia=request.form.get('programme_neolia'),
            adresse=request.form.get('adresse'),
            ville=request.form.get('ville'),
            immatriculation=request.form.get('immatriculation'),
            nombre_logements=int(request.form.get('nombre_logements') or 0),
            exercice_comptable=request.form.get('exercice_comptable'),
            gestionnaire=request.form.get('gestionnaire'),
            comptable=request.form.get('comptable'),
            responsable_secteur=request.form.get('responsable_secteur'),
            commercial=request.form.get('commercial'),
            notaire=request.form.get('notaire'),
            est_active=True
        )
        db.session.add(new_copro)
        db.session.commit()
        flash(f'Copropriété #{numero} ajoutée avec succès !', 'success')
        return redirect(url_for('index'))
    return render_template('new_copropriete.html')

@app.route('/copropriete/<int:copro_id>/edit', methods=['GET', 'POST'])
def edit_copropriete(copro_id):
    copropriete = Copropriete.query.get_or_404(copro_id)

    if request.method == 'POST':
        # Récupère les données du formulaire
        numero = int(request.form.get('numero'))
        nom = request.form.get('nom')
        date_mise_copropriete = parse_date(request.form.get('date_mise_copropriete'))
        programme_neolia = request.form.get('programme_neolia')
        adresse = request.form.get('adresse')
        ville = request.form.get('ville')
        immatriculation = request.form.get('immatriculation')
        nombre_logements = int(request.form.get('nombre_logements') or 0)
        exercice_comptable = request.form.get('exercice_comptable')
        gestionnaire = request.form.get('gestionnaire')
        comptable = request.form.get('comptable')
        responsable_secteur = request.form.get('responsable_secteur')
        commercial = request.form.get('commercial')
        notaire = request.form.get('notaire')
        est_active = 'est_active' in request.form

        # Vérifie si le numéro existe déjà (sauf pour la copropriété actuelle)
        existing = Copropriete.query.filter(
            Copropriete.numero == numero,
            Copropriete.id != copro_id
        ).first()
        if existing:
            flash(f'Le numéro {numero} existe déjà !', 'error')
            return redirect(url_for('edit_copropriete', copro_id=copro_id))

        # Met à jour les champs
        copropriete.numero = numero
        copropriete.nom = nom
        copropriete.date_mise_copropriete = date_mise_copropriete
        copropriete.programme_neolia = programme_neolia
        copropriete.adresse = adresse
        copropriete.ville = ville
        copropriete.immatriculation = immatriculation
        copropriete.nombre_logements = nombre_logements
        copropriete.exercice_comptable = exercice_comptable
        copropriete.gestionnaire = gestionnaire
        copropriete.comptable = comptable
        copropriete.responsable_secteur = responsable_secteur
        copropriete.commercial = commercial
        copropriete.notaire = notaire
        copropriete.est_active = est_active

        db.session.commit()
        flash(f'Copropriété #{numero} modifiée avec succès !', 'success')
        return redirect(url_for('copropriete', copro_id=copro_id))

    # Si GET, affiche le formulaire avec les données actuelles
    return render_template('edit_copropriete.html', copropriete=copropriete)

@app.route('/copropriete/<int:copro_id>/toggle_active', methods=['POST'])
def toggle_active(copro_id):
    copro = Copropriete.query.get_or_404(copro_id)
    copro.est_active = not copro.est_active
    db.session.commit()
    action = "désactivée" if not copro.est_active else "réactivée"
    flash(f'Copropriété #{copro.numero} {action} avec succès !', 'success')
    return redirect(url_for('index'))

# ========== FICHE IMMEUBLE ROUTES ==========
@app.route('/fiche-immeuble/save', methods=['POST'])
def save_fiche_immeuble():
    copro_id = request.form.get('copro_id')
    copropriete = Copropriete.query.get(copro_id)
    if not copropriete:
        flash('Copropriété non trouvée', 'error')
        return redirect(url_for('index'))

    fiche = FicheImmeuble.query.filter_by(copropriete_id=copro_id).first()
    if not fiche:
        fiche = FicheImmeuble(copropriete_id=copro_id)

    fiche.date_construction = parse_date(request.form.get('date_construction'))
    fiche.date_arrete_compte = parse_date(request.form.get('date_arrete_compte'))
    fiche.reference_neolia = request.form.get('reference_neolia')
    fiche.nom_commercial_vente_hlm = request.form.get('nom_commercial_vente_hlm')
    fiche.nom_responsable_secteur = request.form.get('nom_responsable_secteur')
    fiche.nom_notaire = request.form.get('nom_notaire')
    fiche.designation_immeuble = request.form.get('designation_immeuble')
    fiche.iban = request.form.get('iban')
    fiche.ics = request.form.get('ics')
    fiche.lien_dossier = request.form.get('lien_dossier')
    fiche.lien_budget = request.form.get('lien_budget')

    db.session.add(fiche)
    db.session.commit()
    flash('Fiche immeuble sauvegardée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

# ========== CONTRATS ROUTES ==========
@app.route('/contrat/save', methods=['POST'])
def save_contrat():
    contrat_id = request.form.get('contrat_id')
    copro_id = request.form.get('copro_id')

    if contrat_id:
        contrat = Contrat.query.get(contrat_id)
        if not contrat:
            flash('Contrat non trouvé', 'error')
            return redirect(url_for('copropriete', copro_id=copro_id))
    else:
        contrat = Contrat(copropriete_id=copro_id)

    contrat.type_contrat = request.form.get('type_contrat')
    contrat.nature = request.form.get('nature')
    contrat.fournisseur = request.form.get('fournisseur')
    contrat.date_debut = parse_date(request.form.get('date_debut'))
    contrat.date_fin = parse_date(request.form.get('date_fin'))
    contrat.montant_annuel = float(request.form.get('montant_annuel') or 0)

    db.session.add(contrat)
    db.session.commit()
    flash('Contrat sauvegardé avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

@app.route('/contrat/<int:contrat_id>/delete', methods=['POST'])
def delete_contrat(contrat_id):
    contrat = Contrat.query.get_or_404(contrat_id)
    copro_id = contrat.copropriete_id
    db.session.delete(contrat)
    db.session.commit()
    flash('Contrat supprimé avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

# ========== COPROPRIETAIRE ROUTES ==========
@app.route('/save_coproprietaire', methods=['POST'])
def save_coproprietaire():
    _migrer_coproprietaires()
    copro_id = request.form.get('copro_id')
    coproprietaire_id = request.form.get('coproprietaire_id')

    # --- Récupération des données simples ---
    nom = request.form.get('persons[0][nom]') or request.form.get('nom')
    prenom = request.form.get('persons[0][prenom]') or request.form.get('prenom')
    civilite_id = request.form.get('persons[0][civilite_id]') or request.form.get('civilite_id')
    date_acquisition = parse_date(request.form.get('date_acquisition'))
    statut = request.form.get('statut')
    # Compatibilité : la fiche copropriété utilise des radios « statut »
    # (residence/loue/autre), tandis que d'autres formulaires utilisent des
    # champs cachés est_residence_principale/est_loue. On déduit le statut
    # depuis les deux formats.
    if 'est_residence_principale' in request.form:
        est_residence_principale = True
    elif statut == 'residence':
        est_residence_principale = True
    else:
        est_residence_principale = False
    if 'est_loue' in request.form:
        est_loue = True
    elif statut == 'loue':
        est_loue = True
    else:
        est_loue = False
    est_conseil_syndical = 'est_conseil_syndical' in request.form
    date_envoi_mail_accueil = parse_date(request.form.get('date_envoi_mail_accueil'))
    lien_espace_client = request.form.get('lien_espace_client')

    # --- Récupération des infos du locataire ---
    locataire_nom = request.form.get('locataire_nom')
    locataire_civilite_id = request.form.get('locataire_civilite_id')
    locataire_prenom = request.form.get('locataire_prenom')
    locataire_email = request.form.get('locataire_email')
    locataire_telephone = request.form.get('locataire_telephone')

    # --- Validation ---
    if est_residence_principale and est_loue:
        flash("Un copropriétaire ne peut pas être à la fois en résidence principale et avoir un lot loué.", "error")
        return redirect(url_for('copropriete', copro_id=copro_id))

    if est_loue and not locataire_nom:
        flash("Le nom du locataire est obligatoire si le lot est loué.", "error")
        return redirect(url_for('copropriete', copro_id=copro_id))

    # --- Récupération des LOTS (nouveau format) ---
    lots_data = []
    for key in request.form:
        if key.startswith('lots[') and key.endswith('][numero]'):
            index = key.split('[')[1].split(']')[0]
            numero = request.form.get(f'lots[{index}][numero]')
            nature = request.form.get(f'lots[{index}][nature]', '')
            if numero:
                lots_data.append({'numero': numero, 'nature': nature})

    if not lots_data:
        flash("Au moins un lot est obligatoire.", "error")
        return redirect(url_for('copropriete', copro_id=copro_id))

    # --- Récupération des EMAILS ---
    emails_data = []
    principal_email_count = 0
    for key in request.form:
        if key.startswith('emails[') and key.endswith('][value]'):
            index = key.split('[')[1].split(']')[0]
            value = request.form.get(f'emails[{index}][value]')
            principal = request.form.get(f'emails[{index}][principal]') == '1'
            if value:
                if principal:
                    principal_email_count += 1
                emails_data.append({'value': value, 'principal': principal})

    if principal_email_count > 1:
        flash("Un seul email peut être marqué comme principal.", "error")
        return redirect(url_for('copropriete', copro_id=copro_id))

    # --- Récupération des TÉLÉPHONES ---
    telephones_data = []
    principal_phone_count = 0
    for key in request.form:
        if key.startswith('telephones[') and key.endswith('][value]'):
            index = key.split('[')[1].split(']')[0]
            value = request.form.get(f'telephones[{index}][value]')
            principal = request.form.get(f'telephones[{index}][principal]') == '1'
            if value:
                if principal:
                    principal_phone_count += 1
                telephones_data.append({'value': value, 'principal': principal})

    if principal_phone_count > 1:
        flash("Un seul téléphone peut être marqué comme principal.", "error")
        return redirect(url_for('copropriete', copro_id=copro_id))

    # --- Sauvegarde ---
    try:
        if coproprietaire_id:
            cp = Coproprietaire.query.get(coproprietaire_id)
            if not cp:
                flash('Copropriétaire non trouvé', 'error')
                return redirect(url_for('copropriete', copro_id=copro_id))
        else:
            cp = Coproprietaire(copropriete_id=copro_id)

        cp.nom = nom
        cp.prenom = prenom
        cp.date_acquisition = date_acquisition
        cp.est_residence_principale = est_residence_principale
        cp.est_loue = est_loue
        cp.est_conseil_syndical = est_conseil_syndical
        cp.date_envoi_mail_accueil = date_envoi_mail_accueil
        # Conserve le lien espace client existant si le champ n'est pas
        # soumis (le champ a été retiré des formulaires).
        if 'lien_espace_client' in request.form:
            cp.lien_espace_client = lien_espace_client
        cp.locataire_nom = locataire_nom
        cp.locataire_prenom = locataire_prenom
        cp.locataire_email = locataire_email
        cp.civilite_id = civilite_id
        cp.locataire_civilite_id = locataire_civilite_id
        cp.locataire_telephone = locataire_telephone
        cp.adresse = request.form.get('adresse')
        cp.code_postal = request.form.get('code_postal')
        cp.ville = request.form.get('ville')

        cp.lots = [LotCoproprietaire(numero=lot['numero'], nature=lot['nature']) for lot in lots_data]
        cp.emails = [EmailCoproprietaire(value=email['value'], principal=email['principal']) for email in emails_data]
        cp.telephones = [TelephoneCoproprietaire(value=tel['value'], principal=tel['principal']) for tel in telephones_data]

        db.session.add(cp)
        db.session.commit()
        flash('Copropriétaire sauvegardé avec succès !', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f"Erreur lors de la sauvegarde : {str(e)}", "error")
        app.logger.error(f"Erreur: {str(e)}")

    if request.form.get('from') == 'coproprietaires':
        return redirect(url_for('coproprietaires_page'))
    return redirect(url_for('copropriete', copro_id=copro_id))


@app.route('/coproprietaire/<int:coproprietaire_id>/toggle-conseil-syndical', methods=['POST'], endpoint='coproprietaire_toggle_conseil_syndical')
def coproprietaire_toggle_conseil_syndical(coproprietaire_id):
    """Coche/décoche le statut « Membre du Conseil Syndical » d'un
    copropriétaire (case à cocher directe dans le tableau)."""
    _migrer_coproprietaires()
    cp = Coproprietaire.query.get_or_404(coproprietaire_id)
    cp.est_conseil_syndical = not cp.est_conseil_syndical
    db.session.commit()
    flash('Statut Conseil Syndical mis à jour.', 'success')
    return redirect(request.referrer or url_for('coproprietaires_page'))


# ========== AG ROUTES ==========
@app.route('/ag/save', methods=['POST'])
def save_ag():
    copro_id = request.form.get('copro_id')
    ag_id = request.form.get('ag_id')

    if ag_id:
        ag = AssembleeGenerale.query.get(ag_id)
        if not ag:
            flash('AG non trouvée', 'error')
            return redirect(url_for('copropriete', copro_id=copro_id))
    else:
        ag = AssembleeGenerale(copropriete_id=copro_id)

    ag.date = parse_date(request.form.get('date'))
    ag.horaire_debut = request.form.get('horaire_debut')
    ag.lieu = request.form.get('lieu')
    ag.type_ag = request.form.get('type_ag')

    db.session.add(ag)
    db.session.commit()
    flash('AG sauvegardée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

# ========== NOUVELLES ROUTES POUR SUPPRIMER AG ET RESOLUTION ==========

@app.route('/ag/<int:ag_id>/delete', methods=['POST'])
def delete_ag(ag_id):
    ag = AssembleeGenerale.query.get_or_404(ag_id)
    copro_id = ag.copropriete_id
    db.session.delete(ag)
    db.session.commit()
    flash('Assemblée Générale supprimée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

# ========== RESOLUTION ROUTES ==========
@app.route('/resolution/save', methods=['POST'])
def save_resolution():
    copro_id = request.form.get('copro_id')
    resolution_id = request.form.get('resolution_id')

    if resolution_id:
        resolution = ResolutionFuture.query.get(resolution_id)
        if not resolution:
            flash('Résolution non trouvée', 'error')
            return redirect(url_for('copropriete', copro_id=copro_id))
    else:
        resolution = ResolutionFuture(copropriete_id=copro_id)

    resolution.titre = request.form.get('titre')
    resolution.projet = request.form.get('projet')

    db.session.add(resolution)
    db.session.commit()
    flash('Résolution sauvegardée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

@app.route('/resolution/<int:resolution_id>/delete', methods=['POST'])
def delete_resolution(resolution_id):
    resolution = ResolutionFuture.query.get_or_404(resolution_id)
    copro_id = resolution.copropriete_id
    db.session.delete(resolution)
    db.session.commit()
    flash('Résolution supprimée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

# ========== DEMANDES ROUTES ==========
@app.route('/demande/save', methods=['POST'])
def save_demande():
    copro_id = request.form.get('copro_id')
    demande_id = request.form.get('demande_id')

    if demande_id:
        demande = Demande.query.get(demande_id)
        if not demande:
            flash('Demande non trouvée', 'error')
            return redirect(url_for('copropriete', copro_id=copro_id))
    else:
        demande = Demande(copropriete_id=copro_id)

    demande.date = parse_date(request.form.get('date'))
    demande.demandeur = request.form.get('demandeur')
    if demande.demandeur == '__autre__':
        demande.demandeur = request.form.get('demandeur_autre', '').strip()
    demande.type_action = request.form.get('type_action')
    demande.details = request.form.get('details')
    demande.echeance = request.form.get('echeance')
    demande.suivi = request.form.get('suivi', 'a_faire')

    db.session.add(demande)
    db.session.commit()
    flash('Demande sauvegardée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

@app.route('/demande/<int:demande_id>/delete', methods=['POST'])
def delete_demande(demande_id):
    demande = Demande.query.get_or_404(demande_id)
    copro_id = demande.copropriete_id
    db.session.delete(demande)
    db.session.commit()
    flash('Demande supprimée avec succès !', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

@app.route('/demande/<int:demande_id>/suivi', methods=['POST'])
def demande_suivi(demande_id):
    demande = Demande.query.get_or_404(demande_id)
    copro_id = demande.copropriete_id
    suivi = request.form.get('suivi', 'a_faire')
    if suivi not in ('a_faire', 'en_cours', 'fait'):
        suivi = 'a_faire'
    demande.suivi = suivi
    db.session.commit()
    flash('Suivi mis à jour.', 'success')
    return redirect(url_for('copropriete', copro_id=copro_id))

def _migrer_colonnes_coproprietaires():
    """Ajoute les colonnes adresse/code_postal/ville à la table des
    copropriétaires si elles manquent (db.create_all ne modifie pas une
    table existante). Compatible avec les noms de table « coproprietaire »
    et « coproprietaires ».
    """
    from sqlalchemy import inspect, text
    engine = db.engine
    insp = inspect(engine)
    tables = insp.get_table_names()
    table = 'coproprietaire' if 'coproprietaire' in tables else 'coproprietaires'
    if table not in tables:
        return
    cols = {c['name'] for c in insp.get_columns(table)}
    with engine.begin() as conn:
        if 'adresse' not in cols:
            conn.execute(text(f"ALTER TABLE {table} ADD COLUMN adresse VARCHAR(300)"))
        if 'code_postal' not in cols:
            conn.execute(text(f"ALTER TABLE {table} ADD COLUMN code_postal VARCHAR(20)"))
        if 'ville' not in cols:
            conn.execute(text(f"ALTER TABLE {table} ADD COLUMN ville VARCHAR(100)"))


def _importer_coproprietaires_excel():
    """Importe les copropriétaires initiaux (une seule fois, si la table
    est vide). Les données de l'ancien fichier « Copropriétaires.xlsx »
    sont intégrées au logiciel (coproprietaires_data.py) : plus besoin du
    fichier Excel."""
    if Coproprietaire.query.count() > 0:
        return

    copros = {c.numero: c for c in Copropriete.query.all()}
    civilites = {c.libelle: c for c in Civilite.query.all()}

    nb = 0
    for d in COPROPRIETAIRES_EXCEL_DATA:
        copro = copros.get(d['numero_copro'])
        if not copro:
            continue
        civilite = civilites.get(d['civilite']) if d['civilite'] else None
        lots_str = d['lots'] or ''
        cp = Coproprietaire(
            copropriete_id=copro.id,
            civilite_id=civilite.id if civilite else None,
            nom=d['nom'],
            prenom=d['prenom'],
            adresse=d['adresse'],
            code_postal=str(d['code_postal']) if d['code_postal'] else None,
            ville=d['ville'],
            telephone=d['telephone'],
            email=d['email'],
        )
        premier_lot = next((p.strip() for p in lots_str.replace(';', ',').split(',') if p.strip()), None)
        if premier_lot and premier_lot.upper() != 'NC':
            cp.numero_lot = premier_lot
        db.session.add(cp)
        nb += 1
    db.session.commit()
    if nb:
        app.logger.info(f"{nb} copropriétaires importés depuis les données intégrées")


# ========== MAIN ==========
if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        _migrer_colonnes_coproprietaires()

        # Initialiser les civilités si elles n'existent pas
        civilites_data = ["Monsieur", "Madame", "Monsieur et Madame", "Société"]
        for libelle in civilites_data:
            if not Civilite.query.filter_by(libelle=libelle).first():
                civilite = Civilite(libelle=libelle)
                db.session.add(civilite)
        db.session.commit()
        # Initialiser les copropriétés si elles n'existent pas
        if Copropriete.query.count() == 0:
            for data in COPROPRIETES_DATA:
                copro = Copropriete(
                    numero=data["numero"],
                    nom=f"Copropriété {data['numero']}",
                    date_mise_copropriete=parse_date(data["date_mise_copropriete"]),
                    programme_neolia=data["programme_neolia"],
                    adresse=data["adresse"],
                    ville=data["ville"],
                    immatriculation=data["immatriculation"],
                    nombre_logements=data["nombre_logements"],
                    exercice_comptable=data["exercice_comptable"],
                    gestionnaire=data["gestionnaire"],
                    est_active=data["est_active"]
                )
                db.session.add(copro)
            db.session.commit()

        # Ajouter la colonne Conseil Syndical aux bases existantes
        _migrer_coproprietaires()

        # Importer les copropriétaires depuis le fichier Excel (une fois)
        _importer_coproprietaires_excel()

        # Importer les honoraires depuis les données « Honoraire agate » (une fois)
        _importer_honoraires()

        # Importer les visites depuis les données « Visites d'immeuble agate » (une fois)
        _importer_visites()

        # Importer les travaux depuis les données « ADF Travaux AGATE » (une fois)
        _importer_travaux()
